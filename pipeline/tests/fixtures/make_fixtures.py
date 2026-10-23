"""Trim raw downloads under data/raw/ into small test fixtures.

Run from the pipeline directory after refreshing data/raw/:
    python tests/fixtures/make_fixtures.py
Only public web pages and open-data files are used.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import random
import re
import zipfile
from collections import defaultdict
from pathlib import Path

import openpyxl

from lmia_pipeline.normalize import norm_name, split_trade_name
from lmia_pipeline.sources import fmcsa, lfs, vacancies
from lmia_pipeline.sources.matching import is_person_name

HERE = Path(__file__).resolve().parent
RAW = HERE.parents[2] / "data" / "raw"
WAGES_2025_URL = (
    "https://open.canada.ca/data/dataset/adad580f-76b0-4502-bd05-20c125de9116/resource/"
    "9da94d63-b178-4a64-aeb3-b6a3bd721ad2/download/2a71-das-wage2025opendata-esdc-all-19nov2025-vf.csv"
)

ARTICLE_RE = re.compile(r"<article[^>]*id=\"article-\d+\".*?</article>", re.S)
SCRIPT_RE = re.compile(r"<(script|style)\b.*?</\1>", re.S | re.I)


def read_html(src: Path) -> str:
    data = src.read_bytes()
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def trim_search(src: Path, dst: Path, keep: int = 4) -> None:
    html = read_html(src)
    articles = ARTICLE_RE.findall(html)
    first = html.find(articles[0])
    last_end = html.find(articles[-1]) + len(articles[-1])
    trimmed = html[:first] + "\n".join(articles[:keep]) + html[last_end:]
    dst.write_text(SCRIPT_RE.sub("", trimmed), encoding="utf-8")


def trim_detail(src: Path, dst: Path) -> None:
    dst.write_text(SCRIPT_RE.sub("", read_html(src)), encoding="utf-8")


def trim_xlsx(src: Path, dst: Path, keep: int = 40) -> None:
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    out = openpyxl.Workbook()
    ws = out.active
    for row in rows[: keep + 2] + rows[-3:]:
        ws.append(list(row))
    out.save(dst)


def trim_ircc(html_src: Path, json_src: Path, dst_html: Path, dst_json: Path) -> None:
    html = read_html(html_src)
    items = [m.group(0) for m in re.finditer(r"<(\w+)[^>]*id=\"list\d+\"[^>]*>.*?</\1>", html, re.S)]
    dst_html.write_text("<html><body><ol>" + "\n".join(items) + "</ol></body></html>", encoding="utf-8")
    data = json.loads(json_src.read_text(encoding="utf-8"))
    dst_json.write_text(json.dumps({"list": data["list"][:3]}, ensure_ascii=False, indent=1), encoding="utf-8")


def cached(url: str) -> Path:
    """The pipeline's cache file for a URL (http.PoliteClient keys it by sha1)."""
    return RAW / "cache" / f"{hashlib.sha1(url.encode('utf-8')).hexdigest()}.bin"


# Regional figures (44101), provincial only in Manitoba (22101), national only
# (21103) and an occupation published at an annual rate (00010).
WAGE_NOCS = ("NOC_44101", "NOC_22101", "NOC_21103", "NOC_00010")


def trim_wages(src: Path, dst: Path) -> None:
    text = src.read_bytes().decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    keep = [rows[0]] + [
        r for r in rows[1:]
        if r and r[0] in WAGE_NOCS and (r[4] == "ER00" or r[4].startswith("ER46") or r[4] in ("ER10", "ER1010"))
    ]
    with dst.open("w", newline="", encoding="utf-8") as fh:
        csv.writer(fh, lineterminator="\n").writerows(keep)


LFS_GEOS = ("2021A000011124", "2021S05031", "2021S0503505", "2021S050535505", "2021S0503602")
LFS_MONTHS = {f"{y}-{m:02d}" for y in (2024, 2025, 2026) for m in (3, 6, 9, 12)}


def trim_lfs(src: Path, dst: Path) -> None:
    with zipfile.ZipFile(src) as zf:
        text = zf.read("14100459.csv").decode("utf-8-sig")
        meta = zf.read("14100459_MetaData.csv")
    rows = list(csv.reader(io.StringIO(text)))
    keep = [rows[0]] + [
        r for r in rows[1:]
        if r[2] in LFS_GEOS and r[0] in LFS_MONTHS and r[3] in ("Unemployment rate", "Population")
        and r[4] in ("Estimate", "Standard error of estimate")
    ]
    out = io.StringIO()
    csv.writer(out, lineterminator="\n").writerows(keep)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("14100459.csv", out.getvalue().encode("utf-8"))
        zf.writestr("14100459_MetaData.csv", meta[:2000])


GEO_FSAS = {"R2H", "R3C", "R3T", "R0C", "R0G", "R0E", "R1A"}


def _thin(ring: list, limit: int = 120) -> list:
    step = max(1, len(ring) // limit)
    return [[round(x, 4), round(y, 4)] for x, y in ring[::step]]


def trim_geography(dst: Path) -> None:
    from lmia_pipeline.sources.geography import layer_url

    out: dict[str, dict] = {}
    for layer in ("er", "cma", "fsa"):
        data = json.loads(cached(layer_url(layer, "46")).read_bytes())
        feats = data["features"]
        if layer == "fsa":
            feats = [f for f in feats if f["attributes"]["CFSAUID"] in GEO_FSAS]
        for f in feats:
            # Small islands and lakes only add bytes; the tests place points on land.
            rings = sorted(f["geometry"]["rings"], key=len, reverse=True)
            f["geometry"] = {"rings": [_thin(r) for r in rings[:1] + [r for r in rings[1:] if len(r) >= 40]]}
        out[layer] = {"features": feats}
    dst.write_text(json.dumps(out, separators=(",", ":")), encoding="utf-8")


# Winnipeg, its province and Canada, one combined region, and a region with no
# Manitoba tie (Avalon Peninsula); cooks, wage-fixture occupations, a minor
# group (dropped) and the all-occupations total, ten quarters (eight kept).
VACANCY_GEOS = ("2021A000011124", "2021A000246", "2021S05004650", "2021S05001010", "")
VACANCY_COMBINED = "South Central and North Central, Manitoba"
VACANCY_NOCS = ("[63200]", "[44101]", "[22101]", "[4410]", "Total, all occupations")


def trim_vacancies(src: Path, dst: Path) -> None:
    with zipfile.ZipFile(src) as zf:
        meta = zf.read("14100444_MetaData.csv")
        with zf.open("14100444.csv") as fh:
            text = io.TextIOWrapper(fh, encoding="utf-8-sig", newline="")
            reader = csv.reader(text)
            keep = [next(reader)]
            for r in reader:
                if r[0] < "2024-01" or r[2] not in VACANCY_GEOS or (r[2] == "" and r[1] != VACANCY_COMBINED):
                    continue
                if any(r[3].endswith(code) or r[3] == code for code in VACANCY_NOCS):
                    keep.append(r)
    out = io.StringIO()
    csv.writer(out, quoting=csv.QUOTE_ALL, lineterminator="\n").writerows(keep)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("14100444.csv", ("﻿" + out.getvalue()).encode("utf-8"))
        # The cube's title row is enough: it carries the last reference month.
        zf.writestr("14100444_MetaData.csv", meta.replace(b"\r\n", b"\n").split(b"\n\n")[0] + b"\n")


NOC_KEEP = {"6", "63", "632", "6320", "63200", "4", "44", "441", "4410", "44101", "2", "22", "221", "2210", "22101"}


def trim_noc(src: Path, dst: Path) -> None:
    # Some definitions hold line breaks inside quotes; the repository keeps LF only.
    lines = src.read_bytes().decode("utf-8-sig").replace("\r\n", "\n").splitlines(keepends=True)
    rows = list(csv.reader(lines))
    keep = [rows[0]] + [r for r in rows[1:] if len(r) > 2 and r[2] in NOC_KEEP]
    out = io.StringIO()
    csv.writer(out, lineterminator="\n").writerows(keep)
    dst.write_bytes(("﻿" + out.getvalue()).encode("utf-8"))


# -- public records -------------------------------------------------------------
# A fixture holds no row that names a person: besides the stage's own test
# (is_person_name on the legal name), a name that still reads as a person once
# its legal form is stripped ("Warren Gibson Ltd") is left out too.


def _personish(name: str | None) -> bool:
    return bool(name) and (is_person_name(name) or is_person_name(norm_name(name)))


def trim_grants(src: Path, dst: Path) -> None:
    """The first 40 MB of grants.csv (a streamed, cut-off download) -> about 270 rows.

    Bytes are kept as published (read through latin-1, which maps every byte
    to one character and back); long free text is blanked and titles cut to 60
    bytes. Kept: every amended agreement up to 40, then rows with a program
    account after the business number, nine-digit numbers, placeholders,
    Accessibility Standards Canada rows and a few of each recipient type.
    """
    raw = src.read_bytes()
    text = raw[: raw.rfind(b"\n") + 1].decode("latin-1").removeprefix("\u00ef\u00bb\u00bf")
    rows = list(csv.reader(io.StringIO(text, newline="")))
    header, body = rows[0], [r for r in rows[1:] if len(r) == len(rows[0])]
    col = {h: i for i, h in enumerate(header)}

    def utf8(v: str) -> str:
        return v.encode("latin-1").decode("utf-8", errors="replace")

    def cut(v: str, n: int) -> str:
        b = v.encode("latin-1")[:n]
        while b and b[-1] >= 0x80:
            b = b[:-1]
        return b.decode("latin-1").rstrip()

    keep = [r for r in body if r[col["recipient_type"]] != "P" and not any(
        _personish(split_trade_name(utf8(r[col[k]]))[0])
        for k in ("recipient_legal_name", "recipient_operating_name", "research_organization_name"))]
    by_ref: dict[tuple, list] = defaultdict(list)
    for r in keep:
        by_ref[(r[col["owner_org"]], r[col["ref_number"]])].append(r)
    random.seed(7)
    chosen = [r for v in [v for v in by_ref.values() if len(v) > 1][:40] for r in v]
    pool = [r for r in keep if len(by_ref[(r[col["owner_org"]], r[col["ref_number"]])]) == 1]
    bn = col["recipient_business_number"]
    nine = [r for r in pool if len(r[bn].strip()) == 9 and r[bn] != "000000000"]
    types: dict[str, list] = defaultdict(list)
    for r in pool:
        types[r[col["recipient_type"]]].append(r)
    chosen += [r for r in pool if len(r[bn].strip()) >= 15][:25] + random.sample(nine, min(80, len(nine)))
    chosen += [r for r in pool if r[bn].strip() == "000000000"][:10]
    chosen += [r for r in pool if r[col["owner_org"]] == "casdo-ocena"][:20]
    for v in types.values():
        chosen += v[:8]
    seen: set[tuple] = set()
    out_rows = []
    for r in chosen:
        if tuple(r) in seen:
            continue
        seen.add(tuple(r))
        r = list(r)
        for k in ("prog_purpose_en", "prog_purpose_fr", "description_en", "description_fr", "expected_results_en",
                  "expected_results_fr", "additional_information_en", "additional_information_fr"):
            r[col[k]] = ""
        for k in ("agreement_title_en", "agreement_title_fr"):
            r[col[k]] = cut(r[col[k]], 60)
        out_rows.append(r)
    out = io.StringIO(newline="")
    csv.writer(out, lineterminator="\n").writerows([header] + out_rows[:320])
    dst.write_bytes(bytes([0xEF, 0xBB, 0xBF]) + out.getvalue().encode("latin-1"))


def trim_alberta(src: Path, dst: Path) -> None:
    """The Employer Records workbook -> about 300 rows of the Injury sheet, whole accounts only.

    Picked: accounts listed under two names, accounts with a fatality, with
    claims, without a city, numbered companies, then a random rest; never an
    account named as a person or a professional corporation. An empty Order
    sheet stays, so the parser must find its sheet by name.
    """
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    rows = wb["Injury (2020-2024)"].iter_rows(values_only=True)
    header = next(rows)
    order_header = next(wb["Order (2020-2024)"].iter_rows(values_only=True))
    by_key: dict[str, list] = defaultdict(list)
    for row in rows:
        by_key[row[0]].append(row)
    random.seed(11)
    keys = [k for k, v in by_key.items()
            if not _personish(split_trade_name(v[0][1] or "")[0]) and "PROF" not in (v[0][1] or "").upper()]
    pick = [k for k in keys if len({r[7] for r in by_key[k]}) < len(by_key[k])][:4]
    pick += [k for k in keys if any((r[11] or 0) > 0 for r in by_key[k])][:5]
    pick += random.sample([k for k in keys if any((r[9] or 0) > 0 for r in by_key[k])], 12)
    pick += random.sample([k for k in keys if not by_key[k][0][3]], 6)
    pick += random.sample([k for k in keys if (by_key[k][0][1] or "")[:1].isdigit()], 6)
    chosen: list[str] = []
    for k in pick + random.sample(keys, 200):
        if k not in chosen:
            chosen.append(k)
        if sum(len(by_key[c]) for c in chosen) >= 300:
            break
    out = openpyxl.Workbook()
    sheet = out.active
    sheet.title = "Injury (2020-2024)"
    sheet.append(list(header))
    for k in chosen:
        for r in by_key[k]:
            sheet.append(list(r))
    out.create_sheet("Order (2020-2024)").append(list(order_header))
    out.save(dst)


def trim_fmcsa(src: Path, dst: Path) -> None:
    """One API page of 1,500 Canadian carriers (fmcsa.FIELDS, $order=:id) -> 300 records.

    Picked: rated carriers, active ones, records without a postal code, records
    with a DBA name, then a random rest; never an individual or a name that
    reads as a person.
    """
    recs = [r for r in json.loads(src.read_text(encoding="utf-8"))
            if r.get("business_org_desc") != "INDIVIDUAL"
            and not any(_personish(split_trade_name(r.get(k) or "")[0]) for k in ("legal_name", "dba_name"))]
    assert all(set(r) <= set(fmcsa.FIELDS) for r in recs), "a field outside fmcsa.FIELDS was fetched"
    random.seed(5)
    pick = random.sample([r for r in recs if r.get("safety_rating")], 80)
    pick += [r for r in recs if r.get("status_code") == "A"][:60]
    pick += random.sample([r for r in recs if (r.get("phy_zip") or "").upper() == "CANADA"], 20)
    pick += [r for r in recs if r.get("dba_name")][:30]
    seen: set[str] = set()
    out = []
    for r in pick + random.sample(recs, 300):
        if r["dot_number"] not in seen and len(out) < 300:
            seen.add(r["dot_number"])
            out.append(r)
    out.sort(key=lambda r: int(r["dot_number"]))
    dst.write_bytes(json.dumps(out, ensure_ascii=False, indent=0).encode("utf-8"))


def main() -> None:
    trim_search(RAW / "jobbank" / "jb_bc_requested.html", HERE / "jobbank_requested.html")
    trim_search(RAW / "jobbank" / "jb_bc_approved.html", HERE / "jobbank_approved.html")
    trim_detail(RAW / "jobbank" / "jb_50456228.html", HERE / "jobbank_detail_50456228.html")
    trim_detail(RAW / "jobbank" / "jb_50440025.html", HERE / "jobbank_detail_50440025.html")
    trim_xlsx(RAW / "esdc" / "tfwp_2026q1_pos_en.xlsx", HERE / "esdc_2026q1_pos_trim.xlsx")
    trim_xlsx(RAW / "esdc" / "tfwp_2026q1_neg_en.xlsx", HERE / "esdc_2026q1_neg_trim.xlsx")
    trim_ircc(RAW / "ircc" / "noncompliant.html", RAW / "ircc" / "non_compliant.json",
              HERE / "ircc_legend.html", HERE / "ircc_sample.json")
    # The next three read the pipeline's download cache, so run the wages,
    # lfs and geography stages once first.
    trim_wages(cached(WAGES_2025_URL), HERE / "wages_trim.csv")
    trim_lfs(cached(lfs.URL), HERE / "lfs_14100459_trim.zip")
    trim_geography(HERE / "geography_mb.json")
    # These two read the cache the vacancies stage fills.
    trim_vacancies(cached(vacancies.URL), HERE / "vacancies_14100444_trim.zip")
    trim_noc(cached(vacancies.NOC_URLS["en"]), HERE / "noc2021_en_trim.csv")
    trim_noc(cached(vacancies.NOC_URLS["fr"]), HERE / "noc2021_fr_trim.csv")
    # The public records, from files saved by hand once (the stages keep no
    # copy): the first 40 MB of grants.csv, the Employer Records workbook, and
    # one API page of Canadian carriers requested with fmcsa.FIELDS.
    trim_grants(RAW / "grants" / "grants_head.csv", HERE / "grants_trim.csv")
    trim_alberta(RAW / "alberta" / "employer_records_2020_2024.xlsx", HERE / "alberta_employer_records_trim.xlsx")
    trim_fmcsa(RAW / "fmcsa" / "canada_page.json", HERE / "fmcsa_canada_trim.json")
    print("fixtures written to", HERE)


if __name__ == "__main__":
    main()
