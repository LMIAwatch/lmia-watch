"""Alberta Employer Records: the workbook parser on a trimmed copy, the listing cadence and the join.

The fixture is 300 rows of the "Injury (2020-2024)" sheet of the live
workbook (dated 2025-12-16), whole accounts only and none named as a person,
with an empty "Order (2020-2024)" sheet so the parser must pick its sheet by
name. The person names in the synthetic workbook below are invented.
"""

import io
import json

import openpyxl
import pytest

from conftest import FakeResponse
from lmia_pipeline import export_static
from lmia_pipeline.db import get_meta, upsert_employer
from lmia_pipeline.normalize import norm_name
from lmia_pipeline.sources import alberta_records

NOW = "2026-10-22T00:00:00Z"
FIXTURE = "alberta_employer_records_trim.xlsx"
WORKBOOK_URL = ("https://open.alberta.ca/dataset/a2772d8c-48be-4d39-bcf2-dafca456d724/resource/"
                "fb8cd062-039a-4f4f-b012-a4220fdc8a67/download/2024_ohs-employer-record-open-data.xlsx")


def listing(modified: str = "2025-12-16T22:57:20.949112") -> bytes:
    return json.dumps({"result": {"resources": [
        {"id": "bd81bc8b", "format": "DOCX", "created": "2017-10-05T10:44:36",
         "url": "https://open.alberta.ca/dataset/x/resource/bd81bc8b/download/employer-records-open-data-fields.docx"},
        {"id": "fb8cd062-039a-4f4f-b012-a4220fdc8a67", "format": "XLSX", "last_modified": modified, "url": WORKBOOK_URL},
    ]}}).encode()


def every_name(data: bytes) -> set[str]:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    return {norm_name(r[1]) for r in wb["Injury (2020-2024)"].iter_rows(min_row=2, values_only=True) if r[1]}


def test_parse_reads_the_injury_sheet_with_one_row_per_account_year_and_name(fixtures):
    data = (fixtures / FIXTURE).read_bytes()
    rows, counts = alberta_records.parse_workbook(data, every_name(data), "res")
    # "S.Q.I - STORY..." and "S.Q.I. - STORY..." normalize alike: one row a year.
    assert counts == {"read": 300, "persons": 0, "kept": 298}
    by = {(r[0][:8], r[1], r[2]): dict(zip(alberta_records.COLUMNS, r)) for r in rows}
    cargill = by[("814469DD", 2020, "CARGILL LIMITED")]
    assert (cargill["city"], cargill["industry"], cargill["industry_code"]) == ("HIGH RIVER", "MEAT PROCESSING", "10100")
    assert (cargill["person_years"], cargill["lost_time_claims"], cargill["lost_time_rate"]) == (3046.59, 785, 25.77)
    assert (cargill["disabling_injuries"], cargill["disabling_rate"], cargill["fatalities"]) == (843, 27.67, 2)
    assert cargill["cor_year"] is None and cargill["trade_name"] is None and cargill["resource_id"] == "res"
    # One account listed under two names keeps both, with the same figures.
    pine = by[("0B351169", 2022, "PINE CLIFF ENERGY LTD.")]
    certus = by[("0B351169", 2022, "CERTUS OIL AND GAS INC.")]
    assert pine["person_years"] == certus["person_years"] == 29.075
    # A blank city stays blank.
    assert by[("64D9EC7B", 2020, "CHAHAL EXPRESS LINE INC.")]["city"] is None


def synthetic_workbook(rows: list[tuple]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Injury (2021-2025)"
    ws.append(["EMP_KEY", "EMPLOYER_NAME", "TRADE_NAME", "CITY_NAME", "WCB_INDUSTRY_NAME", "WCB_INDUSTRY_CODE",
               "COR_YEAR", "YEAR_NO", "PERSON_YEARS_COUNT", "LOST_TIME_CLAIM_COUNT", "LOST_TIME_CLAIM_RATE",
               "ANNUAL_FATALITY_COUNT", "MOTOR_VEHICLE_FATALITY_COUNT", "WORKPLACE_INCIDENT_FATALITY_COUNT",
               "OCCUPATIONAL_DISEASE_FATALITY_COUNT", "ANNUAL_DISABLING_INJURIES_COUNT", "DISABLING_INJURIES_RATE"])
    for key, name, trade, city, year in rows:
        ws.append([key, name, trade, city, "ROOFING", "42118", 2024, year, 4.2, 1, None, 0, 0, 0, None, 1, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_an_account_of_an_individual_is_left_out_whole_and_a_person_trade_name_is_not_kept():
    data = synthetic_workbook([
        ("K1", "ANA FERREIRA", None, "CALGARY", 2024),
        ("K1", "FERREIRA ROOFING LTD.", None, "CALGARY", 2024),
        ("K2", "CORVID ROOFING LTD.", "OWEN T MERCER", "EDMONTON", 2024),
        ("K3", "LENA OKAFOR, MIRA LAW", None, "RED DEER", 2024),
    ])
    wanted = {"ferreira roofing", "corvid roofing", "ana ferreira", "lena okafor mira law"}
    rows, counts = alberta_records.parse_workbook(data, wanted, "res")
    assert counts == {"read": 4, "persons": 2, "kept": 1}
    kept = dict(zip(alberta_records.COLUMNS, rows[0]))
    assert (kept["employer_name"], kept["trade_name"], kept["cor_year"]) == ("CORVID ROOFING LTD.", None, 2024)


def test_a_workbook_without_the_sheet_or_its_columns_is_refused():
    wb = openpyxl.Workbook()
    wb.active.title = "Order (2020-2024)"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="no Injury sheet"):
        alberta_records.parse_workbook(buf.getvalue(), set(), "res")
    wb = openpyxl.Workbook()
    wb.active.title = "Injury (2020-2024)"
    wb.active.append(["EMP_KEY", "EMPLOYER_NAME"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="lacks columns"):
        alberta_records.parse_workbook(buf.getvalue(), set(), "res")


def seed(conn):
    def ab(name, city=None, postal=None, province="AB"):
        return upsert_employer(conn, name, province, source="esdc_positive", city=city, postal_code=postal, now=NOW)

    return {
        # Name and city agree: exact.
        "cargill": ab("Cargill Limited", city="High River"),
        # Name agrees, the city does not: probable.
        "pine": ab("Pine Cliff Energy Ltd.", city="Edmonton"),
        # The account's other name, with its city: exact, on the same account.
        "certus": ab("Certus Oil and Gas Inc.", city="Calgary"),
        # A numbered company with no city on record: probable.
        "numbered": ab("1776155 Alberta Ltd.", postal="T8V 1A1"),
        # One word and another city: too common for the province alone.
        "katalyst": ab("Katalyst Corporation", city="Red Deer"),
        # The same name outside Alberta: never compared.
        "cargill_bc": ab("Cargill Limited", city="Chilliwack", province="BC"),
    }


def test_load_and_match(conn, fake_client, fixtures):
    ids = seed(conn)
    client = fake_client({alberta_records.LISTING: listing(), WORKBOOK_URL: (fixtures / FIXTURE).read_bytes()})
    assert alberta_records.run(conn, client, now=NOW) == 4
    found = {(r["employer_id"], r["emp_key"][:8]): (r["grade"], r["basis"], r["account_name"])
             for r in conn.execute("SELECT * FROM alberta_matches")}
    assert found == {
        (ids["cargill"], "814469DD"): ("exact", "name_city", "CARGILL LIMITED"),
        (ids["pine"], "0B351169"): ("probable", "name_province", "PINE CLIFF ENERGY LTD."),
        (ids["certus"], "0B351169"): ("exact", "name_city", "CERTUS OIL AND GAS INC."),
        (ids["numbered"], "1D544AFF"): ("probable", "name_province", "1776155 ALBERTA LTD."),
    }
    # Only the names of Alberta employers were kept from the workbook.
    names = {r[0] for r in conn.execute("SELECT DISTINCT employer_name FROM alberta_employer_records")}
    assert names == {"CARGILL LIMITED", "PINE CLIFF ENERGY LTD.", "CERTUS OIL AND GAS INC.", "1776155 ALBERTA LTD.",
                     "KATALYST CORPORATION"}
    detail = conn.execute("SELECT detail FROM sources_log WHERE source = 'alberta_records'").fetchone()[0]
    assert detail == "rows=23 loaded_now=23 matched=4"

    # The profile shows the latest year of each account, under the name that matched.
    out = export_static.profile_alberta(conn, [ids["certus"], ids["pine"], ids["cargill"], ids["katalyst"]])
    assert out[ids["certus"]][0]["employer_name"] == "CERTUS OIL AND GAS INC." and out[ids["certus"]][0]["year"] == 2023
    assert out[ids["pine"]][0]["employer_name"] == "PINE CLIFF ENERGY LTD." and out[ids["pine"]][0]["grade"] == "probable"
    latest = out[ids["cargill"]][0]
    assert (latest["year"], latest["first_year"], latest["person_years"], latest["fatalities"]) == (2024, 2020, 2524.202, 0)
    assert ids["katalyst"] not in out


def test_the_listing_is_checked_monthly_and_the_workbook_fetched_when_it_changes_or_ages(conn, fake_client, fixtures):
    seed(conn)
    data = (fixtures / FIXTURE).read_bytes()
    client = fake_client({alberta_records.LISTING: listing(), WORKBOOK_URL: data})
    alberta_records.run(conn, client, now=NOW)
    urls = lambda: [u for u, _ in client.fake_session.calls]  # noqa: E731
    assert urls() == [alberta_records.LISTING, WORKBOOK_URL]
    # Within 30 days: no request at all, and the matching still runs.
    assert alberta_records.run(conn, client, now="2026-11-10T00:00:00Z") == 4
    assert len(urls()) == 2
    # After 30 days the listing is read; the same file, loaded 40 days ago, is not fetched.
    client.forget(alberta_records.LISTING)  # the runner starts each night with an empty cache
    alberta_records.run(conn, client, now="2026-12-01T00:00:00Z")
    assert urls()[2:] == [alberta_records.LISTING]
    # A new file is fetched.
    changed = fake_client({alberta_records.LISTING: listing("2026-12-20T00:00:00"), WORKBOOK_URL: data})
    changed.forget(alberta_records.LISTING)  # the runner starts each night with an empty cache
    alberta_records.run(conn, changed, now="2027-01-05T00:00:00Z")
    assert [u for u, _ in changed.fake_session.calls] == [alberta_records.LISTING, WORKBOOK_URL]
    assert get_meta(conn, alberta_records.RESOURCE_KEY).endswith("2026-12-20T00:00:00")


def test_a_failed_download_is_recorded_and_the_rows_stay(conn, fake_client, fixtures):
    seed(conn)
    alberta_records.run(conn, fake_client({alberta_records.LISTING: listing(),
                                           WORKBOOK_URL: (fixtures / FIXTURE).read_bytes()}), now=NOW)
    broken = fake_client({alberta_records.LISTING: listing("2026-12-20T00:00:00"), WORKBOOK_URL: FakeResponse(404)})
    broken.forget(alberta_records.LISTING)
    assert alberta_records.run(conn, broken, now="2027-01-05T00:00:00Z") == 4
    row = conn.execute("SELECT ok, detail FROM sources_log WHERE source = 'alberta_records' ORDER BY id DESC").fetchone()
    assert row["ok"] == 0 and "download failed" in row["detail"]
    assert conn.execute("SELECT COUNT(*) FROM alberta_employer_records").fetchone()[0] == 23
